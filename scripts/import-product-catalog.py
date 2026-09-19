import io
import json
import posixpath
import re
import sys
import zipfile
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from PIL import Image, ImageOps


VISIBLE_SHEETS = ["已量产目录", "客户专用型号目录", "20260625新增13款"]
PRIORITY = {name: index for index, name in enumerate(VISIBLE_SHEETS, start=1)}
NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def normalized(value):
    return re.sub(r"[^A-Z0-9]", "", str(value or "").upper())


def resolve(source_path, target):
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(posixpath.dirname(source_path), target))


def relationships(archive, rel_path, source_path):
    root = ET.fromstring(archive.read(rel_path))
    return {
        item.attrib["Id"]: {
            "target": resolve(source_path, item.attrib["Target"]),
            "type": item.attrib.get("Type", ""),
        }
        for item in root.findall("rel:Relationship", NS)
    }


def sheet_paths(archive):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = relationships(archive, "xl/_rels/workbook.xml.rels", "xl/workbook.xml")
    result = {}
    for sheet in workbook.findall("main:sheets/main:sheet", NS):
        rid = sheet.attrib.get(f"{{{NS['r']}}}id")
        if rid in rels:
            result[sheet.attrib["name"]] = rels[rid]["target"]
    return result


def row_images(archive, sheet_path):
    rel_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels")
    if rel_path not in archive.namelist():
        return {}
    sheet_rels = relationships(archive, rel_path, sheet_path)
    drawings = [item["target"] for item in sheet_rels.values() if item["type"].endswith("/drawing")]
    if not drawings:
        return {}
    drawing_path = drawings[0]
    drawing_rel_path = posixpath.join(posixpath.dirname(drawing_path), "_rels", posixpath.basename(drawing_path) + ".rels")
    drawing_rels = relationships(archive, drawing_rel_path, drawing_path)
    root = ET.fromstring(archive.read(drawing_path))
    result = {}
    for anchor in list(root):
        origin = anchor.find("xdr:from", NS)
        blip = anchor.find(".//a:blip", NS)
        if origin is None or blip is None:
            continue
        row_node = origin.find("xdr:row", NS)
        col_node = origin.find("xdr:col", NS)
        if row_node is None or col_node is None or int(col_node.text or -1) != 2:
            continue
        rid = blip.attrib.get(f"{{{NS['r']}}}embed")
        media = drawing_rels.get(rid, {}).get("target")
        if media:
            result.setdefault(int(row_node.text or 0), media)
    return result


def make_thumbnail(raw, destination):
    with Image.open(io.BytesIO(raw)) as image:
        image = ImageOps.exif_transpose(image).convert("RGBA")
        image.thumbnail((220, 150), Image.Resampling.LANCZOS)
        canvas = Image.new("RGB", (220, 150), "white")
        left = (canvas.width - image.width) // 2
        top = (canvas.height - image.height) // 2
        canvas.paste(image, (left, top), image if image.mode == "RGBA" else None)
        canvas.save(destination, "JPEG", quality=76, optimize=True, progressive=True)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("usage: import-product-catalog.py SOURCE.xlsx PROJECT_ROOT")
    source = Path(sys.argv[1])
    project = Path(sys.argv[2])
    image_dir = project / "public" / "assets" / "catalog"
    data_path = project / "data" / "source" / "产品目录60625.json"
    image_dir.mkdir(parents=True, exist_ok=True)
    data_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(source, read_only=True, data_only=True)
    records = {}
    with zipfile.ZipFile(source) as archive:
        paths = sheet_paths(archive)
        for sheet_name in VISIBLE_SHEETS:
            if sheet_name not in workbook.sheetnames or sheet_name not in paths:
                continue
            worksheet = workbook[sheet_name]
            images = row_images(archive, paths[sheet_name])
            for row_number, values in enumerate(worksheet.iter_rows(min_row=4, values_only=True), start=4):
                internal = re.sub(r"\s+", "", str(values[1] or "").strip()) if len(values) > 1 else ""
                key = normalized(internal)
                if not key:
                    continue
                image_source = images.get(row_number - 1)
                image_url = ""
                if image_source and image_source in archive.namelist():
                    image_name = f"{key}.jpg"
                    make_thumbnail(archive.read(image_source), image_dir / image_name)
                    image_url = f"/assets/catalog/{image_name}"
                record = {
                    "internal": internal,
                    "normalizedInternal": key,
                    "commonName": str(values[3] or "").strip() if len(values) > 3 else "",
                    "vehicle": str(values[4] or "").strip() if len(values) > 4 else "",
                    "oe": str(values[5] or "").strip() if len(values) > 5 else "",
                    "note": str(values[7] or "").strip() if sheet_name == "已量产目录" and len(values) > 7 else (str(values[6] or "").strip() if len(values) > 6 else ""),
                    "imageUrl": image_url,
                    "sourceSheet": sheet_name,
                    "sourceRow": row_number,
                    "priority": PRIORITY[sheet_name],
                }
                existing = records.get(key)
                if not existing or record["priority"] >= existing["priority"]:
                    records[key] = record

    output = {
        "filename": source.name,
        "visibleSheets": VISIBLE_SHEETS,
        "records": sorted(records.values(), key=lambda item: item["normalizedInternal"]),
    }
    data_path.write_text(json.dumps(output, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({"records": len(output["records"]), "images": sum(bool(item["imageUrl"]) for item in output["records"]), "data": str(data_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
